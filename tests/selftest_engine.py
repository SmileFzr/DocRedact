#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""DocRedact 全量验收脚本：引擎 + 格式 + 反向 OCR 验证。"""

# --- 统一入口：无论从哪个目录执行，都以仓库根目录为工作目录 ---
import os as _os, sys as _sys
from pathlib import Path as _Path
_ROOT = _Path(__file__).resolve().parent.parent
_os.chdir(_ROOT)
_sys.path.insert(0, str(_ROOT))
_ROOT.joinpath("outputs").mkdir(exist_ok=True)

import io
import os
import time

import fitz
import numpy as np
from PIL import Image

import redact_engine as e

PDF = "outputs/_TEST_sample.pdf"
OUT = "outputs/_TEST_out.pdf"
SHOT = "outputs/_TEST_shot.png"
SHOT_OUT = "outputs/_TEST_shot_out.png"

KW = ["遮光", "NICETOWN", "张伟",
      {"value": r"1[3-9]\d{9}", "regex": True, "label": "手机号"},
      {"value": r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", "regex": True, "label": "邮箱"}]

results = []


def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))
    print(f"{'✅' if cond else '❌'} {name}" + (f" | {detail}" if detail else ""))


def ocr_text_of_page(page, zoom=150 / 72):
    eng, _ = e.get_ocr_engine()
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
    img = np.array(Image.open(io.BytesIO(pix.tobytes("png"))))
    items = e._parse_ocr_result(e._ocr_call(eng, img))
    return "".join(it["text"] for it in items)


def ocr_text_of_image(path):
    eng, _ = e.get_ocr_engine()
    img = np.array(Image.open(path).convert("RGB"))
    items = e._parse_ocr_result(e._ocr_call(eng, img))
    return "".join(it["text"] for it in items)


def main():
    print("引擎:", e.ocr_engine_desc())
    print("=" * 60)

    # ── 1. 扫描与遮罩计数一致 ────────────────────────────
    t0 = time.time()
    det = e.scan_file(PDF, KW, use_ocr=True)
    scan_total = sum(v["text"] + v["ocr"] for v in det.values())
    red_total = e.redact_file(PDF, OUT, KW, use_ocr=True)
    check("扫描计数 = 遮罩计数", scan_total == red_total,
          f"{scan_total} vs {red_total}, {time.time()-t0:.1f}s")
    check("扫描命中数 > 0", scan_total > 0, str(scan_total))

    # ── 2. 文本层关键词彻底移除 ─────────────────────────
    src = fitz.open(PDF)
    red = fitz.open(OUT)
    leaks = []
    for i, page in enumerate(red):
        t = page.get_text()
        for k in ("遮光", "NICETOWN", "张伟", "13812345678", "zhangwei@example.com"):
            if k in t:
                leaks.append((i + 1, k))
    check("输出文本层无残留关键词", not leaks, str(leaks[:6]))

    # ── 3. 相邻文字不被误删 ─────────────────────────────
    t1 = red[0].get_text()
    keep1 = [w for w in ("面料", "性能", "客单价", "主要卖点", "仅内部使用") if w in t1]
    check("第1页相邻词保留", len(keep1) == 5, str(keep1))
    t4 = red[3].get_text()
    check("第4页相邻词保留(申请/配件/售后)",
          all(w in t4 for w in ("申请", "配件", "售后")),
          f"申请={t4.count('申请')} 配件={t4.count('配件')}")

    # ── 4. 反向 OCR 验证：遮罩后的图片页重新识别应找不到关键词 ──
    txt_p2 = ocr_text_of_page(red[1])
    check("纯图片页遮罩后 OCR 无「遮光」", "遮光" not in txt_p2, txt_p2[:40])
    check("纯图片页遮罩后 OCR 无「NICETOWN」",
          "NICETOWN" not in txt_p2.upper(), "")

    # 原图未遮罩时应能识别到（证明验证本身有效）
    txt_src_p2 = ocr_text_of_page(src[1])
    check("对照：原图 OCR 能识别「遮光」", "遮光" in txt_src_p2, "")

    # ── 5. 遮罩区域确实变黑 ─────────────────────────────
    pix = red[1].get_pixmap(matrix=fitz.Matrix(1, 1))
    arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    dark = (arr[:, :, :3].sum(axis=2) < 60).sum()
    check("图片页存在黑色遮罩像素", dark > 500, f"{dark} 像素")

    # ── 6. 图片格式 ─────────────────────────────────────
    shot_src_txt = ocr_text_of_image(SHOT)
    n_img = e.redact_file(SHOT, SHOT_OUT, KW)
    shot_out_txt = ocr_text_of_image(SHOT_OUT)
    check("图片：原图含关键词", "遮光" in shot_src_txt, "")
    check("图片：遮罩后关键词消失",
          all(k not in shot_out_txt for k in ("遮光", "张伟", "13812345678")),
          f"{n_img} 处 | 残留: {[k for k in ('遮光','张伟','13812345678') if k in shot_out_txt]}")

    # ── 7. 纯文本层模式（关 OCR）────────────────────────
    n_no = e.redact_file(PDF, "outputs/_TEST_noocr.pdf", KW, use_ocr=False)
    check("关闭 OCR 仍能遮文本层", n_no > 0, f"{n_no} 处")

    # ── 8. 大小写敏感 ───────────────────────────────────
    cs = e.scan_file(PDF, ["nicetown"], use_ocr=False, case_sensitive=True)
    ci = e.scan_file(PDF, ["nicetown"], use_ocr=False, case_sensitive=False)
    check("大小写敏感开关生效",
          cs["nicetown"]["text"] < ci["nicetown"]["text"],
          f"敏感={cs['nicetown']['text']} 不敏感={ci['nicetown']['text']}")

    # ── 9. Office 三种格式 ──────────────────────────────
    from docx import Document
    doc = Document()
    doc.add_paragraph("NICETOWN 遮光窗帘，联系人张伟，电话 13812345678。")
    doc.add_paragraph("保留这句：正常的业务描述。")
    doc.save("outputs/_TEST_o.docx")
    n = e.redact_file("outputs/_TEST_o.docx", "outputs/_TEST_o_out.docx",
                      ["NICETOWN", "遮光", "张伟", {"value": r"1[3-9]\d{9}", "regex": True, "label": "手机号"}])
    d2 = Document("outputs/_TEST_o_out.docx")
    t = d2.paragraphs[0].text
    check("DOCX 遮罩", all(k not in t for k in ("NICETOWN", "遮光", "张伟", "13812345678"))
          and "正常的业务描述" in d2.paragraphs[1].text, f"{n} 处 | {t}")

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "NICETOWN 遮光窗帘报表"
    ws["A2"] = "保留这一格"
    wb.save("outputs/_TEST_o.xlsx")
    e.redact_file("outputs/_TEST_o.xlsx", "outputs/_TEST_o_out.xlsx", ["NICETOWN", "遮光"])
    wb2 = openpyxl.load_workbook("outputs/_TEST_o_out.xlsx")
    v1, v2 = wb2.active["A1"].value, wb2.active["A2"].value
    check("XLSX 遮罩", "NICETOWN" not in v1 and "遮光" not in v1 and "报表" in v1 and v2 == "保留这一格", v1)

    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "NICETOWN 遮光方案"
    prs.save("outputs/_TEST_o.pptx")
    e.redact_file("outputs/_TEST_o.pptx", "outputs/_TEST_o_out.pptx", ["NICETOWN", "遮光"])
    p2 = Presentation("outputs/_TEST_o_out.pptx")
    tt = p2.slides[0].shapes.title.text
    check("PPTX 遮罩", "NICETOWN" not in tt and "遮光" not in tt and "方案" in tt, tt)

    # ── 10. 取消 ────────────────────────────────────────
    import threading
    ev = threading.Event()
    ev.set()
    cancelled = False
    try:
        e.redact_file(PDF, "outputs/_TEST_cancel.pdf", KW, use_ocr=True, cancel_event=ev)
    except e.Cancelled:
        cancelled = True
    check("取消机制生效", cancelled)

    # ── 11. 输出可正常打开 ──────────────────────────────
    check("输出 PDF 页数一致", len(red) == len(src), f"{len(src)} -> {len(red)}")
    red.close()
    src.close()

    print("=" * 60)
    ok = sum(1 for _, c, _ in results if c)
    print(f"通过 {ok}/{len(results)}")
    if ok != len(results):
        print("失败项:", [n for n, c, _ in results if not c])


if __name__ == "__main__":
    main()
