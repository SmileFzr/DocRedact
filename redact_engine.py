#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文档隐私遮罩工具 - 核心遮罩引擎  v3.2
支持: PDF（文本层+OCR层）、图片、Word、Excel、PPT

设计要点:
  1. 每一页 rawdict 只解析一次，所有关键词复用（性能）
  2. OCR 每页最多跑一次，且只在「无文本层」或「图片占比够大」时触发
  3. 支持普通关键词 + 正则表达式，统一 Pattern 抽象
  4. 文本层/OCR 层命中互相去重，计数不重复
  5. 全流程支持取消（cancel_event）
  6. OCR 默认 PP-OCRv5 + 逐字符坐标，遮罩框贴合字形，不漏字也不越界
"""

import io
import re
import logging
import threading
from pathlib import Path

import fitz  # PyMuPDF
from PIL import Image, ImageDraw
import numpy as np

log = logging.getLogger(__name__)

MASK_CHAR = "█"


class Cancelled(Exception):
    """任务被用户取消。"""


# ══════════════════════════════════════════════════════
#  OCR 引擎（加锁懒加载单例，避免多线程重复初始化）
# ══════════════════════════════════════════════════════
_ocr_instance = None
_ocr_error = None
_ocr_lock = threading.Lock()
_ocr_desc = ""


def _new_rapidocr():
    """
    优先 PP-OCRv5（mobile）：
      - 中英混排、小字号识别率显著高于 v4（实测 NICETOWN 14/14 vs 6/14）
      - 推理速度约快一倍
    失败则回退 v4 / 旧版 rapidocr_onnxruntime，保证任何环境都能跑。
    """
    try:
        from rapidocr import RapidOCR, OCRVersion, ModelType
        try:
            eng = RapidOCR(params={
                "Det.ocr_version": OCRVersion.PPOCRV5,
                "Det.model_type": ModelType.MOBILE,
                "Rec.ocr_version": OCRVersion.PPOCRV5,
                "Rec.model_type": ModelType.MOBILE,
                "Global.log_level": "error",
            })
            return eng, "RapidOCR v3 + PP-OCRv5 (mobile)"
        except Exception as e:
            log.warning(f"PP-OCRv5 加载失败，回退 PP-OCRv4: {e}")
            eng = RapidOCR(params={"Global.log_level": "error"})
            return eng, "RapidOCR v3 + PP-OCRv4"
    except ImportError:
        pass

    try:
        from rapidocr_onnxruntime import RapidOCR
        return RapidOCR(), "RapidOCR (onnxruntime 旧版)"
    except ImportError as e:
        raise RuntimeError(f"未安装 OCR 依赖: {e}")


def get_ocr_engine():
    """懒加载 OCR 引擎，返回 (engine, error)。"""
    global _ocr_instance, _ocr_error, _ocr_desc
    if _ocr_instance is not None or _ocr_error is not None:
        return _ocr_instance, _ocr_error

    with _ocr_lock:
        if _ocr_instance is not None or _ocr_error is not None:
            return _ocr_instance, _ocr_error

        try:
            _ocr_instance, _ocr_desc = _new_rapidocr()
            log.info(f"OCR 引擎加载成功: {_ocr_desc}")
        except Exception as e:
            _ocr_error = f"OCR 引擎初始化失败: {e}"
            log.warning(_ocr_error)

    return _ocr_instance, _ocr_error


def ocr_engine_desc():
    """引擎描述，用于前端/健康检查展示。"""
    get_ocr_engine()
    return _ocr_desc or (_ocr_error or "未加载")


def ocr_available():
    engine, _ = get_ocr_engine()
    return engine is not None


# ══════════════════════════════════════════════════════
#  匹配规则抽象（普通关键词 / 正则）
# ══════════════════════════════════════════════════════
class Pattern:
    """一个待遮罩的匹配规则。"""

    __slots__ = ("value", "regex", "label", "key")

    def __init__(self, value, regex=False, label=None):
        self.value = str(value)
        self.regex = bool(regex)
        self.label = str(label) if label else self.value
        self.key = self.label

    def to_dict(self):
        return {"value": self.value, "regex": self.regex, "label": self.label}

    def __repr__(self):
        return f"<Pattern {self.key!r} regex={self.regex}>"


def normalize_patterns(keywords):
    """把前端传来的 str / dict 混合列表统一成 Pattern 列表。"""
    out = []
    seen = set()
    for item in keywords or []:
        if isinstance(item, Pattern):
            p = item
        elif isinstance(item, dict):
            value = item.get("value") or item.get("keyword") or ""
            p = Pattern(value, item.get("regex", False), item.get("label"))
        else:
            p = Pattern(item)

        if not p.value:
            continue
        if p.key in seen:
            continue

        if p.regex:
            try:
                re.compile(p.value)
            except re.error as e:
                raise ValueError(f"正则表达式「{p.key}」有语法错误: {e}")

        seen.add(p.key)
        out.append(p)
    return out


def iter_matches(text, pat, case_sensitive=False):
    """返回 pat 在 text 中所有匹配的 (start, end)，支持重叠匹配。"""
    if not text:
        return []

    if pat.regex:
        flags = 0 if case_sensitive else re.IGNORECASE
        try:
            return [(m.start(), m.end()) for m in re.finditer(pat.value, text, flags)
                    if m.end() > m.start()]
        except re.error:
            return []

    haystack = text if case_sensitive else text.lower()
    needle = pat.value if case_sensitive else pat.value.lower()
    if not needle:
        return []

    spans = []
    start = 0
    while True:
        i = haystack.find(needle, start)
        if i == -1:
            break
        spans.append((i, i + len(needle)))
        start = i + 1  # 允许重叠匹配
    return spans


def count_matches(text, patterns, case_sensitive=False):
    return sum(len(iter_matches(text, p, case_sensitive)) for p in patterns)


# ══════════════════════════════════════════════════════
#  PDF：文本层（字符级，精确到关键词本身）
# ══════════════════════════════════════════════════════
def page_char_spans(page):
    """
    一次性从页面提取字符级数据，返回 [(span_text, chars), ...]。
    必须用 rawdict —— dict 模式不包含 chars。
    """
    spans = []
    try:
        try:
            blocks = page.get_text("rawdict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]
        except Exception:
            blocks = page.get_text("rawdict")["blocks"]
    except Exception as e:
        log.debug(f"rawdict 提取失败: {e}")
        return spans

    for block in blocks:
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                chars = span.get("chars") or []
                if not chars:
                    continue
                text = "".join(ch.get("c", "") for ch in chars)
                if text:
                    spans.append((text, chars))
    return spans


def _rects_from_spans(spans, pat, case_sensitive=False, pad_x=0.0):
    """
    在已提取的字符级 spans 中匹配，返回精确 Rect 列表。

    pad_x 必须为 0：PyMuPDF 删除文字时是「矩形与字符框有任何相交就删」，
    一旦外扩，就会把相邻字符（如「申请遮光配件」里的「请」「配」）一起删掉。
    """
    rects = []
    for text, chars in spans:
        n_chars = len(chars)
        for s, e in iter_matches(text, pat, case_sensitive):
            if s < 0 or e > n_chars or e <= s:
                continue
            window = chars[s:e]
            try:
                x0 = min(c["bbox"][0] for c in window)
                y0 = min(c["bbox"][1] for c in window)
                x1 = max(c["bbox"][2] for c in window)
                y1 = max(c["bbox"][3] for c in window)
            except Exception:
                continue
            rect = fitz.Rect(x0 - pad_x, y0, x1 + pad_x, y1)
            if rect.width > 0.5 and rect.height > 0.5:
                rects.append(rect)
    return rects


def find_keyword_rects_text_layer(page, keyword, case_sensitive=False):
    """兼容旧接口：单关键词文本层查找。"""
    pats = normalize_patterns([keyword])
    if not pats:
        return []
    spans = page_char_spans(page)
    return _rects_from_spans(spans, pats[0], case_sensitive)


# ══════════════════════════════════════════════════════
#  OCR：渲染一次 → 识别一次 → 匹配全部规则
# ══════════════════════════════════════════════════════
def _align_char_boxes(text, word_result):
    """
    把 RapidOCR 的 word_results（逐字符 [字符, 置信度, 四点框]）
    对齐成与 text 一一对应的 bbox 列表（对不上的位置为 None）。

    注意：OCR 返回的文本里有空格，但逐字符结果通常不含空格，
    直接比较长度会整行失败 —— 这里用 difflib 做容错对齐，
    空格等插入字符留 None，由调用方对这些位置退化为比例估算。
    """
    import difflib

    if not text or not word_result:
        return None

    # 1) 展开逐字符框（多字符条目按宽度均分）
    chars, jboxes = [], []
    for item in word_result:
        try:
            ch, pts = item[0], item[2]
            xs = [p[0] for p in pts]
            ys = [p[1] for p in pts]
            x0, x1 = min(xs), max(xs)
            y0, y1 = min(ys), max(ys)
        except Exception:
            return None
        if not ch or x1 <= x0:
            continue
        chars.append(ch)
        if len(ch) == 1:
            jboxes.append((x0, y0, x1, y1))
        else:
            step = (x1 - x0) / len(ch)
            for k in range(len(ch)):
                jboxes.append((x0 + step * k, y0, x0 + step * (k + 1), y1))

    joined = "".join(chars)
    if not joined:
        return None
    if joined == text:
        return jboxes

    # 2) 容错对齐：文本 ↔ 逐字符序列
    sm = difflib.SequenceMatcher(None, text, joined, autojunk=False)
    boxes = [None] * len(text)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for k in range(i2 - i1):
                boxes[i1 + k] = jboxes[j1 + k]
        elif tag == "replace" and (j2 > j1):
            span, js = i2 - i1, j2 - j1
            for k in range(span):
                boxes[i1 + k] = jboxes[j1 + min(js - 1, int(k * js / span))]
        # insert / delete 位置保持 None

    hit = sum(1 for b in boxes if b)
    # 对齐率太低说明识别结果与框严重错位，整行退化为比例估算更安全
    if hit < len(text) * 0.5:
        return None
    return boxes


def _ocr_call(engine, img):
    """统一调用 OCR，尽量带上逐字符框；老引擎不支持时自动降级。"""
    try:
        return engine(img, return_word_box=True)
    except TypeError:
        return engine(img)
    except Exception as e:
        log.debug(f"带 word_box 调用失败，改用普通调用: {e}")
        return engine(img)


def _parse_ocr_result(result):
    """
    兼容新旧 RapidOCR API，统一返回
    [{"box": 四点框, "text": 文本, "score": 置信度, "chars": [bbox...] | None}, ...]
    """
    items = []
    if result is None:
        return items

    # 新版 RapidOCR（v3）：RapidOCROutput 对象
    if hasattr(result, "boxes") and hasattr(result, "txts"):
        boxes = result.boxes if result.boxes is not None else []
        txts = result.txts if result.txts is not None else []
        scores = getattr(result, "scores", None) or [1.0] * len(txts)
        word_results = getattr(result, "word_results", None) or []
        for idx, (box, text, score) in enumerate(zip(boxes, txts, scores)):
            if not text:
                continue
            chars = _align_char_boxes(text, word_results[idx]) \
                if idx < len(word_results) else None
            items.append({"box": box, "text": text, "score": float(score),
                          "chars": chars})
        return items

    # 旧版：([box, text, score], ...) 或 tuple 包裹
    raw = result[0] if isinstance(result, tuple) and result else result
    if isinstance(raw, list):
        for line in raw:
            if len(line) >= 2 and line[1]:
                items.append({"box": line[0], "text": line[1],
                              "score": float(line[2]) if len(line) > 2 else 1.0,
                              "chars": None})
    return items


def _rects_from_ocr_items(items, patterns, case_sensitive, scale, pad=1.5):
    """
    OCR 结果 → {pattern.key: [fitz.Rect, ...]}

    优先用逐字符 bbox：命中范围直接取具体字符的框，遮得准、也不越界。
    只有在字符框缺失（老模型/对齐失败）时，才退回「按行宽比例估算」，
    并按旧策略在两侧留 0.8 个字符宽兜底，避免露半边字。
    """
    out = {p.key: [] for p in patterns}
    for item in items:
        box = item["box"]
        text = item["text"]
        chars = item.get("chars")

        xs = [pt[0] for pt in box]
        ys = [pt[1] for pt in box]
        x_min, x_max = min(xs), max(xs)
        y_min, y_max = min(ys), max(ys)
        line_len = len(text)
        if line_len <= 0:
            continue

        char_w = (x_max - x_min) / line_len
        line_h = y_max - y_min

        for pat in patterns:
            for s, e in iter_matches(text, pat, case_sensitive):
                win = chars[s:e] if (chars is not None and e <= len(chars)) else None
                if win and all(win):
                    wx0 = min(b[0] for b in win)
                    wy0 = min(b[1] for b in win)
                    wx1 = max(b[2] for b in win)
                    wy1 = max(b[3] for b in win)
                    # 字框本身就是字形边界，只需一丁点余量盖住抗锯齿边缘
                    px = max((wx1 - wx0) * 0.06, 0.8)
                    py = max((wy1 - wy0) * 0.08, 0.8)
                    rect = fitz.Rect(
                        (wx0 - px) * scale - pad * 0.5,
                        (wy0 - py) * scale,
                        (wx1 + px) * scale + pad * 0.5,
                        (wy1 + py) * scale,
                    )
                else:
                    margin_x = char_w * 0.8
                    margin_y = max(line_h * 0.15, pad)
                    kw_x0 = max(x_min + s * char_w - margin_x,
                                x_min - char_w * 0.2)
                    kw_x1 = min(x_min + e * char_w + margin_x,
                                x_max + char_w * 0.2)
                    rect = fitz.Rect(
                        kw_x0 * scale - pad,
                        (y_min - margin_y) * scale,
                        kw_x1 * scale + pad,
                        (y_max + margin_y) * scale,
                    )

                if rect.width > 0.5 and rect.height > 0.5:
                    out[pat.key].append(rect)
    return out


def ocr_page_rects(page, patterns, engine, case_sensitive=False, dpi=150,
                   max_side=6000):
    """
    渲染 PDF 页面 → OCR 一次 → 返回 {key: [Rect...]}

    超大页面（如 5000pt 宽的海报/长图）按 max_side 限制缩放，
    否则 150dpi 会渲染出上亿像素，内存直接爆掉。
    """
    zoom = dpi / 72.0
    try:
        longest = max(page.rect.width, page.rect.height)
        if longest > 0 and longest * zoom > max_side:
            zoom = max_side / longest
    except Exception:
        pass

    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    items = _parse_ocr_result(_ocr_call(engine, np.array(img)))
    return _rects_from_ocr_items(items, patterns, case_sensitive, 1.0 / zoom)


def ocr_image_rects(img, patterns, engine, case_sensitive=False, max_side=6000):
    """
    对 PIL 图片 OCR 一次 → {key: [(x0,y0,x1,y1) 像素坐标...]}
    超大图先等比缩小再识别，坐标再映射回原尺寸。
    """
    scale_back = 1.0
    if max(img.size) > max_side:
        k = max_side / max(img.size)
        img = img.resize((max(1, int(img.width * k)), max(1, int(img.height * k))),
                         Image.LANCZOS)
        scale_back = 1.0 / k

    items = _parse_ocr_result(_ocr_call(engine, np.array(img)))
    mapped = _rects_from_ocr_items(items, patterns, case_sensitive, scale=1.0, pad=2)

    return {k: [(r.x0 * scale_back, r.y0 * scale_back,
                 r.x1 * scale_back, r.y1 * scale_back) for r in v]
            for k, v in mapped.items()}


def _page_needs_ocr(page, spans):
    """
    是否值得对整页跑 OCR：
      - 没有文本层 → 必须 OCR
      - 有图片且图片面积占页面 > 2% → OCR（图表/截图里可能有字）
      - 纯文字页 → 跳过（速度提升一个数量级）
    """
    if not any(t.strip() for t, _ in spans):
        return True
    try:
        infos = page.get_image_info()
    except Exception:
        return True
    if not infos:
        return False
    try:
        page_area = abs(page.rect) or 1.0
        img_area = sum(abs(fitz.Rect(i["bbox"])) for i in infos if i.get("bbox"))
        return (img_area / page_area) > 0.02
    except Exception:
        return True


# ══════════════════════════════════════════════════════
#  矩形去重
# ══════════════════════════════════════════════════════
def _overlap_ratio(a, b):
    """交叠面积 / 较小矩形面积。"""
    inter = a & b
    if inter.is_empty:
        return 0.0
    small = min(a.get_area(), b.get_area())
    return inter.get_area() / small if small > 0 else 0.0


def _dedupe_rects(rects, thresh=0.5):
    """去掉高度重叠的矩形（文本层与 OCR 层可能命中同一处）。"""
    kept = []
    for r in rects:
        if any(_overlap_ratio(r, k) > thresh for k in kept):
            continue
        kept.append(r)
    return kept


def _filter_overlapping(rects, against, thresh=0.5):
    if not against:
        return rects
    return [r for r in rects if not any(_overlap_ratio(r, k) > thresh for k in against)]


def _check_cancel(cancel_event):
    if cancel_event is not None and cancel_event.is_set():
        raise Cancelled("任务已取消")


# ══════════════════════════════════════════════════════
#  PDF 整体处理
# ══════════════════════════════════════════════════════
def pdf_scan(src_path, patterns, use_ocr=True, case_sensitive=False,
             progress_cb=None, cancel_event=None):
    """扫描 PDF，返回 {key: {"text": n, "ocr": n, "pages": {...}}}"""
    ocr_engine, _ = get_ocr_engine() if use_ocr else (None, None)
    details = {p.key: {"text": 0, "ocr": 0, "pages": {}} for p in patterns}
    if not details:
        return details

    doc = fitz.open(src_path)
    try:
        total_pages = len(doc)
        for page_num, page in enumerate(doc):
            _check_cancel(cancel_event)
            if progress_cb:
                progress_cb(page_num + 1, total_pages,
                            f"扫描第 {page_num + 1}/{total_pages} 页")

            spans = page_char_spans(page)
            text_rects = {p.key: _rects_from_spans(spans, p, case_sensitive)
                          for p in patterns}

            page_hits = {p.key: {"text": len(text_rects[p.key]), "ocr": 0}
                         for p in patterns}
            for p in patterns:
                details[p.key]["text"] += page_hits[p.key]["text"]

            if ocr_engine is not None and _page_needs_ocr(page, spans):
                try:
                    ocr_map = ocr_page_rects(page, patterns, ocr_engine, case_sensitive)
                except Cancelled:
                    raise
                except Exception as e:
                    log.error(f"第{page_num + 1}页 OCR 扫描失败: {e}")
                    ocr_map = {}

                for p in patterns:
                    # 与文本层重合的 OCR 结果不算新命中，避免重复计数
                    only = _filter_overlapping(ocr_map.get(p.key, []),
                                               text_rects[p.key])
                    page_hits[p.key]["ocr"] = len(only)
                    details[p.key]["ocr"] += len(only)

            for key, hit in page_hits.items():
                if hit["text"] or hit["ocr"]:
                    details[key]["pages"][page_num + 1] = hit

            # 单页大文件 OCR 很慢，页末再检查一次，保证取消能及时生效
            _check_cancel(cancel_event)
    finally:
        doc.close()

    return details


def pdf_redact(src_path, out_path, patterns, use_ocr=True, case_sensitive=False,
               progress_cb=None, cancel_event=None):
    """PDF 精确遮罩，返回遮罩块数。"""
    ocr_engine, _ = get_ocr_engine() if use_ocr else (None, None)

    doc = fitz.open(src_path)
    total = 0
    try:
        total_pages = len(doc)
        for page_num, page in enumerate(doc):
            _check_cancel(cancel_event)
            if progress_cb:
                progress_cb(page_num + 1, total_pages,
                            f"遮罩第 {page_num + 1}/{total_pages} 页")

            spans = page_char_spans(page)
            page_rects = []
            for p in patterns:
                page_rects.extend(_rects_from_spans(spans, p, case_sensitive))

            if ocr_engine is not None and _page_needs_ocr(page, spans):
                try:
                    ocr_map = ocr_page_rects(page, patterns, ocr_engine, case_sensitive)
                    for rects in ocr_map.values():
                        page_rects.extend(rects)
                except Cancelled:
                    raise
                except Exception as e:
                    log.error(f"第{page_num + 1}页 OCR 遮罩失败: {e}")

            page_rects = _dedupe_rects(page_rects)
            for rect in page_rects:
                page.add_redact_annot(rect, fill=(0, 0, 0))
            total += len(page_rects)

            if page_rects:
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS)

            _check_cancel(cancel_event)

        _check_cancel(cancel_event)
        if progress_cb:
            progress_cb(total_pages, total_pages, "保存中...")
        doc.save(out_path, garbage=3, deflate=True)
    finally:
        doc.close()

    return total


# ══════════════════════════════════════════════════════
#  图片处理
# ══════════════════════════════════════════════════════
def image_redact(src_path, out_path, patterns, case_sensitive=False,
                 progress_cb=None, cancel_event=None):
    """图片遮罩：OCR 找关键词 → 画黑块。返回遮罩数。"""
    ocr_engine, err = get_ocr_engine()
    if ocr_engine is None:
        raise RuntimeError(f"图片遮罩需要 OCR 引擎: {err}")

    if progress_cb:
        progress_cb(0, 1, "正在识别图片文字...")
    _check_cancel(cancel_event)

    img = Image.open(src_path).convert("RGB")
    rect_map = ocr_image_rects(img, patterns, ocr_engine, case_sensitive)

    draw = ImageDraw.Draw(img)
    total = 0
    for rects in rect_map.values():
        for (x0, y0, x1, y1) in rects:
            draw.rectangle([x0, y0, x1, y1], fill=(0, 0, 0))
            total += 1

    if progress_cb:
        progress_cb(1, 1, "保存中...")
    img.save(out_path)
    return total


def image_scan(src_path, patterns, case_sensitive=False, progress_cb=None,
               cancel_event=None):
    """图片扫描（仅计数）。"""
    ocr_engine, err = get_ocr_engine()
    if ocr_engine is None:
        raise RuntimeError(f"图片扫描需要 OCR 引擎: {err}")

    if progress_cb:
        progress_cb(0, 1, "正在识别图片文字...")
    _check_cancel(cancel_event)

    img = Image.open(src_path).convert("RGB")
    rect_map = ocr_image_rects(img, patterns, ocr_engine, case_sensitive)
    return {k: {"text": 0, "ocr": len(v),
                "pages": {1: {"text": 0, "ocr": len(v)}} if v else {}}
            for k, v in rect_map.items()}


# ══════════════════════════════════════════════════════
#  Office 文档：关键词替换为 █ 实心块
# ══════════════════════════════════════════════════════
def _mask_text(text, patterns, case_sensitive=False, max_mask=20):
    """返回 (替换后文本, 命中次数)。同一位置只替换一次。"""
    if not text:
        return text, 0

    hits = set()
    for pat in patterns:
        for s, e in iter_matches(text, pat, case_sensitive):
            hits.add((s, e))

    if not hits:
        return text, 0

    # 合并重叠区间，避免重复替换导致错位
    merged = []
    for s, e in sorted(hits):
        if merged and s <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], e))
        else:
            merged.append((s, e))

    out = text
    for s, e in reversed(merged):
        out = out[:s] + MASK_CHAR * min(e - s, max_mask) + out[e:]
    return out, len(hits)


def _docx_paragraphs(doc):
    """遍历 Word 所有可能含文字的段落（正文/表格/页眉页脚）。"""
    for p in doc.paragraphs:
        yield p

    def walk_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield p
                    for p in walk_tables(cell.tables):
                        yield p

    for p in walk_tables(doc.tables):
        yield p

    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            try:
                if part is None or part.is_linked_to_previous:
                    continue
                for p in part.paragraphs:
                    yield p
                for p in walk_tables(part.tables):
                    yield p
            except Exception:
                continue


def docx_redact(src_path, out_path, patterns, case_sensitive=False,
                progress_cb=None, cancel_event=None):
    """
    Word 遮罩。
    优先按 run 替换（保留原有混排格式）；只有当关键词跨 run 时，
    才退化成「整段合并到首个 run」以保证不漏遮。
    """
    from docx import Document

    doc = Document(src_path)
    total = 0

    for p in _docx_paragraphs(doc):
        _check_cancel(cancel_event)
        original = p.text
        if not original:
            continue
        expected = count_matches(original, patterns, case_sensitive)
        if expected == 0:
            continue

        hit = 0
        for run in p.runs:
            if not run.text:
                continue
            masked, n = _mask_text(run.text, patterns, case_sensitive)
            if n:
                run.text = masked
                hit += n

        if hit < expected:
            # 跨 run 命中：整段替换（格式有损，但保证遮住）
            masked, n = _mask_text(original, patterns, case_sensitive)
            if p.runs:
                p.runs[0].text = masked
                for run in p.runs[1:]:
                    run.text = ""
            else:
                p.text = masked
            hit = n

        total += hit

    if progress_cb:
        progress_cb(1, 1, "保存中...")
    doc.save(out_path)
    return total


def docx_scan(src_path, patterns, case_sensitive=False, progress_cb=None,
              cancel_event=None):
    """Word 扫描（仅计数）。"""
    from docx import Document

    doc = Document(src_path)
    details = {p.key: {"text": 0, "ocr": 0, "pages": {}} for p in patterns}
    page = 1
    for p in _docx_paragraphs(doc):
        text = p.text
        if not text:
            continue
        for pat in patterns:
            n = len(iter_matches(text, pat, case_sensitive))
            if n:
                details[pat.key]["text"] += n
                details[pat.key]["pages"].setdefault(page, {"text": 0, "ocr": 0})
                details[pat.key]["pages"][page]["text"] += n
    return details


def xlsx_redact(src_path, out_path, patterns, case_sensitive=False,
                progress_cb=None, cancel_event=None):
    """Excel 遮罩：所有工作表单元格中的关键词替换为 █。"""
    import openpyxl

    wb = openpyxl.load_workbook(src_path)
    total = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            _check_cancel(cancel_event)
            for cell in row:
                if isinstance(cell.value, str):
                    masked, n = _mask_text(cell.value, patterns, case_sensitive)
                    if n:
                        cell.value = masked
                        total += n

    if progress_cb:
        progress_cb(1, 1, "保存中...")
    wb.save(out_path)
    return total


def xlsx_scan(src_path, patterns, case_sensitive=False, progress_cb=None,
              cancel_event=None):
    """Excel 扫描（仅计数），按工作表记页码。"""
    import openpyxl

    wb = openpyxl.load_workbook(src_path, read_only=True, data_only=True)
    details = {p.key: {"text": 0, "ocr": 0, "pages": {}} for p in patterns}
    try:
        for idx, ws in enumerate(wb.worksheets, start=1):
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    for pat in patterns:
                        n = len(iter_matches(value, pat, case_sensitive))
                        if n:
                            details[pat.key]["text"] += n
                            details[pat.key]["pages"].setdefault(
                                idx, {"text": 0, "ocr": 0})
                            details[pat.key]["pages"][idx]["text"] += n
    finally:
        wb.close()
    return details


def _iter_pptx_shapes(shapes):
    """递归遍历形状（含组合形状）。"""
    for shape in shapes:
        yield shape
        try:
            if getattr(shape, "shapes", None) is not None:
                yield from _iter_pptx_shapes(shape.shapes)
        except Exception:
            continue


def pptx_redact(src_path, out_path, patterns, case_sensitive=False,
                progress_cb=None, cancel_event=None):
    """PPT 遮罩：所有文本框/表格/备注中的关键词替换为 █。"""
    from pptx import Presentation

    prs = Presentation(src_path)
    total = 0

    def mask_frame(tf):
        nonlocal total
        for p in tf.paragraphs:
            _check_cancel(cancel_event)
            original = "".join(r.text for r in p.runs)
            if not original:
                continue
            masked, n = _mask_text(original, patterns, case_sensitive)
            if n == 0:
                continue
            total += n
            if p.runs:
                p.runs[0].text = masked
                for r in p.runs[1:]:
                    r.text = ""

    for slide in prs.slides:
        if progress_cb:
            progress_cb(1, 1, "处理幻灯片...")
        for shape in _iter_pptx_shapes(slide.shapes):
            try:
                if shape.has_text_frame:
                    mask_frame(shape.text_frame)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            mask_frame(cell.text_frame)
            except Exception as e:
                log.debug(f"跳过形状 {getattr(shape, 'name', '?')}: {e}")
        try:
            if slide.has_notes_slide:
                for shape in slide.notes_slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        mask_frame(shape.text_frame)
        except Exception:
            pass

    if progress_cb:
        progress_cb(1, 1, "保存中...")
    prs.save(out_path)
    return total


def pptx_scan(src_path, patterns, case_sensitive=False, progress_cb=None,
              cancel_event=None):
    """PPT 扫描（仅计数），按幻灯片记页码。"""
    from pptx import Presentation

    prs = Presentation(src_path)
    details = {p.key: {"text": 0, "ocr": 0, "pages": {}} for p in patterns}

    def count_frame(tf, slide_no):
        for p in tf.paragraphs:
            text = "".join(r.text for r in p.runs)
            if not text:
                continue
            for pat in patterns:
                n = len(iter_matches(text, pat, case_sensitive))
                if n:
                    details[pat.key]["text"] += n
                    details[pat.key]["pages"].setdefault(
                        slide_no, {"text": 0, "ocr": 0})
                    details[pat.key]["pages"][slide_no]["text"] += n

    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape in _iter_pptx_shapes(slide.shapes):
            try:
                if shape.has_text_frame:
                    count_frame(shape.text_frame, slide_no)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            count_frame(cell.text_frame, slide_no)
            except Exception:
                continue
        try:
            if slide.has_notes_slide:
                for shape in slide.notes_slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        count_frame(shape.text_frame, slide_no)
        except Exception:
            pass
    return details


# ══════════════════════════════════════════════════════
#  预览渲染
# ══════════════════════════════════════════════════════
def render_page_png(path, page_no=1, dpi=110, max_side=2200):
    """把 PDF 的某一页渲染成 PNG（字节流），返回 (png_bytes, 总页数)。"""
    doc = fitz.open(path)
    try:
        pages = len(doc)
        if pages == 0:
            raise ValueError("PDF 没有页面")
        page_no = max(1, min(int(page_no or 1), pages))
        page = doc[page_no - 1]

        zoom = dpi / 72.0
        w, h = page.rect.width * zoom, page.rect.height * zoom
        if max(w, h) > max_side:
            zoom *= max_side / max(w, h)

        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        return pix.tobytes("png"), pages
    finally:
        doc.close()


def page_count(path):
    """PDF 返回页数，图片返回 1，Office 返回 0。"""
    ext = Path(path).suffix.lower()
    if ext in PDF_EXTS:
        try:
            doc = fitz.open(path)
            try:
                return len(doc)
            finally:
                doc.close()
        except Exception:
            return 0
    if ext in IMAGE_EXTS:
        return 1
    return 0


# ══════════════════════════════════════════════════════
#  预设隐私规则
# ══════════════════════════════════════════════════════
PRESETS = [
    {"id": "phone_cn", "label": "手机号",
     "pattern": r"1[3-9]\d{9}", "regex": True,
     "desc": "中国大陆 11 位手机号"},
    {"id": "email", "label": "邮箱",
     "pattern": r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", "regex": True,
     "desc": "xxx@xxx.com 形式"},
    {"id": "idcard", "label": "身份证号",
     "pattern": r"[1-9]\d{5}(?:19|20)\d{2}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])\d{3}[\dXx]", "regex": True,
     "desc": "18 位居民身份证"},
    {"id": "bankcard", "label": "银行卡号",
     "pattern": r"\d{16,19}", "regex": True,
     "desc": "16~19 位连续数字"},
    {"id": "ipv4", "label": "IP 地址",
     "pattern": r"(?:\d{1,3}\.){3}\d{1,3}", "regex": True,
     "desc": "IPv4 地址"},
    {"id": "date", "label": "日期",
     "pattern": r"20\d{2}[-/年]\d{1,2}[-/月]\d{1,2}日?", "regex": True,
     "desc": "2024-01-01 等格式"},
]


# ══════════════════════════════════════════════════════
#  统一入口
# ══════════════════════════════════════════════════════
SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".pptx",
    ".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif"
}

PDF_EXTS = {".pdf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif"}
DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx"}
PPTX_EXTS = {".pptx"}


def detect_kind(path):
    ext = Path(path).suffix.lower()
    if ext in PDF_EXTS:
        return "pdf"
    if ext in IMAGE_EXTS:
        return "image"
    if ext in DOCX_EXTS:
        return "docx"
    if ext in XLSX_EXTS:
        return "xlsx"
    if ext in PPTX_EXTS:
        return "pptx"
    return "unknown"


def scan_file(src_path, keywords, use_ocr=True, case_sensitive=False,
              progress_cb=None, cancel_event=None):
    """统一扫描入口。keywords 支持 str / dict 混合列表。"""
    patterns = normalize_patterns(keywords)
    if not patterns:
        raise ValueError("没有有效的关键词")

    kind = detect_kind(src_path)
    if kind == "pdf":
        return pdf_scan(src_path, patterns, use_ocr, case_sensitive,
                        progress_cb, cancel_event)
    if kind == "image":
        return image_scan(src_path, patterns, case_sensitive,
                          progress_cb, cancel_event)
    if kind == "docx":
        return docx_scan(src_path, patterns, case_sensitive,
                         progress_cb, cancel_event)
    if kind == "xlsx":
        return xlsx_scan(src_path, patterns, case_sensitive,
                         progress_cb, cancel_event)
    if kind == "pptx":
        return pptx_scan(src_path, patterns, case_sensitive,
                         progress_cb, cancel_event)
    raise ValueError(f"不支持的文件格式: {Path(src_path).suffix}")


def redact_file(src_path, out_path, keywords, use_ocr=True, case_sensitive=False,
                progress_cb=None, cancel_event=None):
    """统一遮罩入口，返回遮罩命中数。"""
    patterns = normalize_patterns(keywords)
    if not patterns:
        raise ValueError("没有有效的关键词")

    kind = detect_kind(src_path)
    if kind == "pdf":
        return pdf_redact(src_path, out_path, patterns, use_ocr, case_sensitive,
                          progress_cb, cancel_event)
    if kind == "image":
        return image_redact(src_path, out_path, patterns, case_sensitive,
                            progress_cb, cancel_event)
    if kind == "docx":
        return docx_redact(src_path, out_path, patterns, case_sensitive,
                           progress_cb, cancel_event)
    if kind == "xlsx":
        return xlsx_redact(src_path, out_path, patterns, case_sensitive,
                           progress_cb, cancel_event)
    if kind == "pptx":
        return pptx_redact(src_path, out_path, patterns, case_sensitive,
                           progress_cb, cancel_event)
    raise ValueError(f"不支持的文件格式: {Path(src_path).suffix}")
